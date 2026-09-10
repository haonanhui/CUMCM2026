"""Pipeline smoke test — 30-second end-to-end sanity check for contest day.

Creates a synthetic problem fixture in a temp sandbox, then runs:
  ingest -> gates -> novelty_gate -> frozen_numbers -> inject_numbers -> verify

Usage:
    python code/pipeline/tests/smoke_test.py

Exit 0 = all stages behave; non-zero = pipeline broken, fix BEFORE the contest.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

LIB = Path(__file__).resolve().parents[1] / "lib"
sys.path.insert(0, str(LIB))
PY = sys.executable


def sh(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run([PY, *args], capture_output=True, text=True,
                          encoding="utf-8", errors="replace")


def make_fixture(root: Path) -> Path:
    """Synthetic 'official' problem dir: one docx + one xlsx + one csv."""
    import docx
    from openpyxl import Workbook

    d = root / "fixture_official"
    d.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    doc.add_paragraph("问题1：计算 {{demo}} 微构体导通概率，要求给出方法与结果。")
    doc.save(d / "题面.docx")
    wb = Workbook()
    ws = wb.active
    ws.append(["x", "y"])
    for i in range(5):
        ws.append([i, i * 2])
    wb.save(d / "附件.xlsx")
    (d / "extra.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    return d


def main() -> int:
    from common import PROJECT_ROOT

    failures = []

    def check(name: str, proc: subprocess.CompletedProcess, expect_rc: int = 0):
        ok = proc.returncode == expect_rc
        print(f"[{'PASS' if ok else 'FAIL'}] {name} (rc={proc.returncode})")
        if not ok:
            failures.append(name)
            print(proc.stdout[-800:], proc.stderr[-800:])

    with tempfile.TemporaryDirectory(prefix="cumcm_smoke_") as td:
        tmp = Path(td)
        fixture = make_fixture(tmp)
        cfg = tmp / "contest.yaml"
        cfg.write_text(
            f"contest: smoke\nproblem: T\nofficial_dir: \"{fixture.as_posix()}\"\n"
            f"deadline: \"2030-01-01 00:00\"\nmode: auto\n",
            encoding="utf-8",
        )
        work = PROJECT_ROOT / "data" / "contest" / "smoke" / "T"
        try:
            check("ingest", sh(str(LIB / "ingest.py"), "--config", str(cfg)))
            assert (work / "00_ingest" / "manifest.json").exists(), "manifest missing"
            assert (work / "00_ingest" / "problem.md").read_text(encoding="utf-8").count("导通概率"), "docx text not extracted"
            csvs = list((work / "00_ingest" / "data").glob("*.csv"))
            assert len(csvs) == 2, f"expected 2 csvs, got {[c.name for c in csvs]}"

            check("gates", sh(str(LIB / "gates.py"), "--config", str(cfg)))

            deltas = work / "03_model" / "deltas.json"
            deltas.parent.mkdir(parents=True, exist_ok=True)
            deltas.write_text(
                '{"Q1": {"baseline": "蒙特卡洛", "baseline_precedent_freq": "high",'
                ' "deltas": [{"layer": "b", "claim": "耦合修正", "evidence": "x"}]}}',
                encoding="utf-8",
            )
            check("novelty_gate", sh(str(LIB / "novelty_gate.py"), "--input", str(deltas), "--config", str(cfg)))

            fi = work / "04_solve" / "frozen_input.json"
            fi.parent.mkdir(parents=True, exist_ok=True)
            fi.write_text('{"demo_p": {"value": 0.732, "source_file": "s.py", "source_line": 1}}', encoding="utf-8")
            check("freeze", sh(str(LIB / "frozen_numbers.py"), "freeze", "--config", str(cfg), "--input", str(fi)))

            tex = work / "paper_test.tex"
            tex.write_text("概率为 {{NUM:demo_p}}。", encoding="utf-8")
            check("inject", sh(str(LIB / "inject_numbers.py"), "--config", str(cfg), "--tex", str(tex)))
            assert "0.732" in tex.read_text(encoding="utf-8"), "injection missing"

            solver = tmp / "q_t.py"
            solver.write_text(
                "from pathlib import Path\n"
                f"out = Path(r'{(work / '04_solve' / 'q_t').as_posix()}')\n"
                "out.mkdir(parents=True, exist_ok=True)\n"
                "(out / 'result.csv').write_text('k,v\\n1,0.5\\n2,0.7\\n', encoding='utf-8')\n",
                encoding="utf-8",
            )
            sh(str(solver))  # first run
            check("verify", sh(str(LIB / "verify.py"), "--config", str(cfg), "--scripts", str(solver)))

            # check command must WARN (exit 1) on unfrozen literal
            tex.write_text("概率为 0.7315。", encoding="utf-8")
            check("frozen-check catches literal", sh(
                str(LIB / "frozen_numbers.py"), "check", "--config", str(cfg), "--tex", str(tex)), expect_rc=1)
        finally:
            import shutil

            shutil.rmtree(PROJECT_ROOT / "data" / "contest" / "smoke", ignore_errors=True)

    print(f"--- smoke: {len(failures)} failure(s) ---")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
